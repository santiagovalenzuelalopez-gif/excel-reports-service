import io
from datetime import datetime
from pathlib import Path
from zoneinfo import ZoneInfo

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import create_engine, text

from app.core.config import Settings
from app.main import app
from app.services.clients import SqliteClientResolver
from app.services.container import get_service, get_verifier
from app.services.identity import StaticIdentityVerifier
from app.services.service import ReportService

AUTH = {"Authorization": "Bearer test-token"}
TZ = ZoneInfo("America/Bogota")

USERS_DDL = """CREATE TABLE users (id INTEGER PRIMARY KEY, name TEXT, login TEXT, doc_type TEXT, doc_number TEXT,
    email TEXT, address TEXT, phone TEXT, country TEXT, department TEXT, city TEXT, kind TEXT, status TEXT)"""
AUDIT_DDL = """CREATE TABLE audit_log (id INTEGER PRIMARY KEY, created_at INTEGER, ip TEXT, user_id INTEGER,
    login TEXT, entity TEXT, entity_2 TEXT, entity_id INTEGER, action TEXT, detail TEXT)"""


def epoch(year, month, day, hour=0, minute=0, second=0) -> int:
    """Epoch de una hora LOCAL de Bogotá (UTC-5): así los tests no dependen del reloj de la máquina."""
    return int(datetime(year, month, day, hour, minute, second, tzinfo=TZ).timestamp())


def make_client_db(path: Path, users=(), audit=()) -> None:
    engine = create_engine(f"sqlite:///{path}")
    with engine.begin() as conn:
        conn.execute(text(USERS_DDL))
        conn.execute(text(AUDIT_DDL))
        for user in users:
            conn.execute(
                text("INSERT INTO users (id, name, login, email, phone, status) VALUES (:id, :name, :login, :email, :phone, :status)"),
                {"email": None, "phone": None, "status": "activo", **user},
            )
        for row in audit:
            conn.execute(
                text("INSERT INTO audit_log (id, created_at, login, action, detail) VALUES (:id, :created_at, :login, :action, :detail)"),
                {"login": "ana", "action": "LOGIN", "detail": "", **row},
            )
    engine.dispose()


def read_sheet(content: bytes, *, read_only=False):
    workbook = load_workbook(io.BytesIO(content), read_only=read_only)
    return workbook[workbook.sheetnames[0]]


@pytest.fixture
def env(tmp_path: Path):
    db_dir = tmp_path / "db"
    db_dir.mkdir()
    settings = Settings(
        sqlite_dir=str(db_dir), demo_clients="acme,beta", static_tokens="test-token",
        report_timezone="America/Bogota", max_rows=1000, max_range_days=31, fetch_chunk_size=50,
    )
    return settings, db_dir


@pytest.fixture
def client(env):
    settings, db_dir = env
    service = ReportService(SqliteClientResolver(settings.sqlite_dir, ["acme", "beta"]), settings)
    app.dependency_overrides[get_service] = lambda: service
    app.dependency_overrides[get_verifier] = lambda: StaticIdentityVerifier(["test-token"])
    with TestClient(app) as test_client:
        yield test_client
    app.dependency_overrides.clear()
