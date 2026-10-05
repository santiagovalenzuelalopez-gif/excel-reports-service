import tracemalloc
from decimal import Decimal

import pytest
from openpyxl import load_workbook
from sqlalchemy import create_engine, text

from app.core.config import Settings
from app.services.clients import mysql_url
from app.services.export import _clean_value, _col_letter, generate_workbook
from app.services.identity import InvalidIdentity, StaticIdentityVerifier
from app.services.reports import USERS, ReportContext
from tests.conftest import TZ, USERS_DDL


def test_mysql_url_escapes_credentials():
    url = mysql_url({"db_user": "u@x", "db_password": "p:/ss@rd", "host": "db.internal", "port": None, "db_name": "d"})
    rendered = url.render_as_string(hide_password=False)
    assert rendered == "mysql+pymysql://u%40x:p%3A%2Fss%40rd@db.internal:3306/d"
    assert url.host == "db.internal"  # la '@' de la contraseña no desplazó el host


def test_clean_value():
    assert _clean_value("a\x00b") == "ab"
    assert _clean_value(Decimal("1.5")) == 1.5
    assert _clean_value(b"hola") == "hola"
    assert _clean_value(None) is None


@pytest.mark.parametrize(("index", "letters"), [(1, "A"), (26, "Z"), (27, "AA"), (52, "AZ"), (53, "BA")])
def test_col_letter(index, letters):
    assert _col_letter(index) == letters


def test_static_verifier():
    verifier = StaticIdentityVerifier(["ok"])
    assert verifier.verify("ok")["sub"]
    with pytest.raises(InvalidIdentity):
        verifier.verify("nope")


def _export_peak(tmp_path, name: str, row_count: int) -> tuple[int, int]:
    """(filas exportadas, pico de memoria de Python en bytes) para una tabla de ``row_count`` filas."""
    engine = create_engine(f"sqlite:///{tmp_path / name}")
    with engine.begin() as conn:
        conn.execute(text(USERS_DDL))
        conn.execute(
            text("INSERT INTO users (id, name, login, email, address) VALUES (:id, :n, :l, :e, :a)"),
            [{"id": i, "n": f"Usuario {i}", "l": f"user{i}", "e": f"u{i}@example.com", "a": "calle " * 20} for i in range(row_count)],
        )
    settings = Settings(max_rows=100_000, fetch_chunk_size=300, spool_max_bytes=64 * 1024)

    tracemalloc.start()
    handle, rows = generate_workbook(engine, USERS, {"letter": "Todos"}, ReportContext(TZ), settings)
    _, peak = tracemalloc.get_traced_memory()
    tracemalloc.stop()

    # el archivo es válido y completo (con un spool de 64 KB se escribió a disco, no en RAM)
    sheet = load_workbook(handle, read_only=True).worksheets[0]
    assert sum(1 for _ in sheet.iter_rows()) == row_count + 1  # + encabezado
    handle.close()
    return rows, peak


def test_memory_is_bounded_by_chunk_size_not_by_table_size(tmp_path):
    """Triplicar la tabla no debe triplicar la memoria: el pico depende del bloque, no del total.
    (Con un DataFrame crecería linealmente: ~3x.)"""
    rows_small, peak_small = _export_peak(tmp_path, "small.db", 3_000)
    rows_big, peak_big = _export_peak(tmp_path, "big.db", 9_000)

    assert (rows_small, rows_big) == (3_000, 9_000)
    assert peak_big < 1.5 * peak_small, f"{peak_small / 1e6:.1f} MB -> {peak_big / 1e6:.1f} MB"
    assert peak_big < 20 * 1024 * 1024


def test_aborted_export_does_not_leave_open_workbook(tmp_path):
    from app.services.export import ReportTooLarge

    engine = create_engine(f"sqlite:///{tmp_path / 'x.db'}")
    with engine.begin() as conn:
        conn.execute(text(USERS_DDL))
        conn.execute(text("INSERT INTO users (id, login) VALUES (:id, 'a')"), [{"id": i} for i in range(50)])
    settings = Settings(max_rows=10, fetch_chunk_size=5)
    with pytest.raises(ReportTooLarge):
        generate_workbook(engine, USERS, {"letter": "Todos"}, ReportContext(TZ), settings)
